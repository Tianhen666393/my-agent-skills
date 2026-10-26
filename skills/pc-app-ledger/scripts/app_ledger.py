import os
import sys
import json
import argparse
import datetime
import hashlib
import shutil
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

def get_documents_dir():
    if sys.platform == 'win32':
        import ctypes.wintypes
        CSIDL_PERSONAL = 5
        SHGFP_TYPE_CURRENT = 0
        buf = ctypes.create_unicode_buffer(ctypes.wintypes.MAX_PATH)
        ctypes.windll.shell32.SHGetFolderPathW(None, CSIDL_PERSONAL, None, SHGFP_TYPE_CURRENT, buf)
        if buf.value:
            return buf.value
    return os.path.expanduser('~/Documents')

def get_ledger_path(device='天选5pro'):
    docs = get_documents_dir()
    return os.path.join(docs, 'my-pc-ledger', f'应用安装清单-{device}.xlsx')

def create_backup(excel_path, reason='修改台账前备份'):
    if not os.path.exists(excel_path):
        return None
    drive = os.path.splitdrive(os.path.abspath(excel_path))[0] or 'C:'
    today = datetime.datetime.now().strftime('%Y-%m-%d')
    path_hash = hashlib.sha1(os.path.abspath(excel_path).lower().encode('utf-8')).hexdigest()[:16]
    ts = datetime.datetime.now().strftime('%Y%m%d-%H%M%S-%f')[:21]
    backup_dir = os.path.join(f"{drive}\\", '.agent_backups', today, f"应用台账-{path_hash}", f"备份-{ts}")
    os.makedirs(backup_dir, exist_ok=True)
    
    target_file = os.path.join(backup_dir, os.path.basename(excel_path))
    shutil.copy2(excel_path, target_file)
    
    # 备份说明.md
    readme_path = os.path.join(backup_dir, '备份说明.md')
    with open(readme_path, 'w', encoding='utf-8') as f:
        f.write(f"# 应用台账备份说明\n\n"
                f"- 原文件完整路径：{os.path.abspath(excel_path)}\n"
                f"- 备份文件：{os.path.basename(excel_path)}\n"
                f"- 备份时间：{datetime.datetime.now(datetime.timezone.utc).isoformat()}\n"
                f"- 执行工具：pc-app-ledger\n"
                f"- 备份原因：{reason}\n"
                f"- 恢复方法：关闭占用该 Excel 的软件，将本目录下的原样副本复制覆盖回原路径。\n")
        
    meta_path = os.path.join(backup_dir, 'backup.json')
    with open(meta_path, 'w', encoding='utf-8') as f:
        json.dump({
            "sourcePath": os.path.abspath(excel_path),
            "createdAt": datetime.datetime.now(datetime.timezone.utc).isoformat(),
            "reason": reason,
            "backupFile": os.path.basename(excel_path)
        }, f, ensure_ascii=False, indent=2)
        
    return backup_dir

def load_records(excel_path):
    if not os.path.exists(excel_path):
        return []
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb.active
    records = []
    for row in ws.iter_rows(values_only=True):
        idx = str(row[0] or '').strip()
        if not idx or '序号' in idx or '应用安装清单' in idx or '截至' in idx:
            continue
        try:
            seq = int(idx)
        except ValueError:
            continue
        if len(row) >= 7:
            records.append({
                'seq': seq,
                'name': str(row[1] or '').strip(),
                'category': str(row[2] or '').strip(),
                'purpose': str(row[3] or '').strip(),
                'version': str(row[4] or '').strip(),
                'path': str(row[5] or '').strip(),
                'notes': str(row[6] or '').strip()
            })
        else:
            records.append({
                'seq': seq,
                'name': str(row[1] or '').strip(),
                'category': '常用应用与工具',
                'purpose': str(row[2] or '').strip(),
                'version': str(row[3] or '').strip(),
                'path': str(row[4] or '').strip(),
                'notes': str(row[5] or '').strip()
            })
    return records

def save_records(excel_path, device, records):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = device
    ws.views.sheetView[0].showGridLines = True
    
    today_str = datetime.datetime.now().strftime('%Y-%m-%d')
    headers = ['序号', '软件名称', '类别', '主要用途', '版本', '安装绝对路径', '备注']
    
    # Title
    ws.merge_cells('A1:G1')
    c_title = ws['A1']
    c_title.value = f'{device} 本机应用安装清单'
    c_title.font = Font(name='Microsoft YaHei', size=16, bold=True, color='17365D')
    c_title.alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[1].height = 32
    
    # Subtitle
    ws.merge_cells('A2:G2')
    c_sub = ws['A2']
    c_sub.value = f'设备：{device} ｜ 截至 {today_str} ｜ 共 {len(records)} 项 ｜ 统一维护个人应用安装路径与备忘'
    c_sub.font = Font(name='Microsoft YaHei', size=10, color='595959')
    c_sub.alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[2].height = 26
    
    # Headers
    h_fill = PatternFill(start_color='17365D', end_color='17365D', fill_type='solid')
    h_font = Font(name='Microsoft YaHei', size=11, bold=True, color='FFFFFF')
    h_align = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[3].height = 30
    for col_idx, h in enumerate(headers, 1):
        c = ws.cell(row=3, column=col_idx, value=h)
        c.fill = h_fill
        c.font = h_font
        c.alignment = h_align
        
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    data_font = Font(name='Microsoft YaHei', size=11)
    align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    
    for row_idx, r in enumerate(records, 4):
        ws.row_dimensions[row_idx].height = 42
        r['seq'] = row_idx - 3
        # Clean control characters that break openpyxl XML export
        clean_name = openpyxl.cell.cell.ILLEGAL_CHARACTERS_RE.sub('', str(r['name'] or ''))
        clean_cat = openpyxl.cell.cell.ILLEGAL_CHARACTERS_RE.sub('', str(r.get('category', '常用应用与工具') or ''))
        clean_purpose = openpyxl.cell.cell.ILLEGAL_CHARACTERS_RE.sub('', str(r['purpose'] or ''))
        clean_version = openpyxl.cell.cell.ILLEGAL_CHARACTERS_RE.sub('', str(r['version'] or ''))
        clean_path = openpyxl.cell.cell.ILLEGAL_CHARACTERS_RE.sub('', str(r['path'] or ''))
        clean_notes = openpyxl.cell.cell.ILLEGAL_CHARACTERS_RE.sub('', str(r['notes'] or ''))
        row_vals = [r['seq'], clean_name, clean_cat, clean_purpose, clean_version, clean_path, clean_notes]
        for col_idx, val in enumerate(row_vals, 1):
            c = ws.cell(row=row_idx, column=col_idx, value=val)
            c.font = data_font
            c.border = thin_border
            c.alignment = align_center if col_idx in [1, 3, 5] else align_left
            
    widths = [8, 24, 22, 28, 14, 65, 36]
    for col_idx, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = w
        
    ws.freeze_panes = 'A4'
    ws.auto_filter.ref = f'A3:G{max(4, len(records)+3)}'
    
    os.makedirs(os.path.dirname(excel_path), exist_ok=True)
    wb.save(excel_path)

def main():
    parser = argparse.ArgumentParser(description="PC App Ledger CLI")
    subparsers = parser.add_subparsers(dest="command", required=True)
    
    # check
    p_check = subparsers.add_parser("check", help="巡检应用安装路径有效性")
    p_check.add_argument("--device", default="天选5pro")
    
    # upsert
    p_upsert = subparsers.add_parser("upsert", help="新增或更新应用记录")
    p_upsert.add_argument("--device", default="天选5pro")
    p_upsert.add_argument("--request", help="包含 entries 列表的 JSON 文件路径")
    p_upsert.add_argument("--apply", action="store_true", help="写入生效")
    
    # remove
    p_remove = subparsers.add_parser("remove", help="移除应用记录")
    p_remove.add_argument("--device", default="天选5pro")
    p_remove.add_argument("--name", required=True, help="要移除的软件名称")
    p_remove.add_argument("--apply", action="store_true", help="写入生效")
    
    # sync
    p_sync = subparsers.add_parser("sync", help="刷新 README 看板")
    
    args = parser.parse_args()
    excel_path = get_ledger_path(args.device if hasattr(args, 'device') else '天选5pro')
    
    if args.command == "check":
        records = load_records(excel_path)
        print(f"检查设备 [{args.device}] 应用台账（共 {len(records)} 项）：")
        for r in records:
            p = r['path']
            exists = os.path.exists(p) if p else False
            status = "已核实有效" if exists else ("路径未填写" if not p else "路径不存在/失效")
            print(f"- [{r['seq']:2d}] {r['name']:25} | {status} | {p}")
            
    elif args.command == "upsert":
        if not args.request or not os.path.exists(args.request):
            print("错误：请提供有效的 --request JSON 文件")
            sys.exit(1)
        with open(args.request, 'r', encoding='utf-8') as f:
            req_data = json.load(f)
        entries = req_data.get('entries', [])
        records = load_records(excel_path)
        
        for entry in entries:
            name = entry.get('name', '').strip()
            if not name:
                continue
            idx = next((i for i, r in enumerate(records) if r['name'].lower() == name.lower()), -1)
            cat = entry.get('category', '').strip()
            if not cat and idx >= 0:
                cat = records[idx].get('category', '常用应用与工具')
            elif not cat:
                cat = '常用应用与工具'

            item = {
                'name': name,
                'category': cat,
                'purpose': entry.get('purpose', '').strip(),
                'version': entry.get('version', '未提供').strip(),
                'path': entry.get('path', '').strip(),
                'notes': entry.get('notes', '').strip()
            }
            if idx >= 0:
                records[idx].update(item)
                print(f"更新已有应用：{name}")
            else:
                records.append(item)
                print(f"新增应用：{name}")
                
        print(f"处理后总项数：{len(records)}（apply={args.apply}）")
        if args.apply:
            b_dir = create_backup(excel_path, "upsert 应用台账")
            if b_dir:
                print(f"已生成安全备份：{b_dir}")
            save_records(excel_path, args.device, records)
            print(f"成功保存至：{excel_path}")
            
    elif args.command == "remove":
        records = load_records(excel_path)
        idx = next((i for i, r in enumerate(records) if r['name'].lower() == args.name.lower()), -1)
        if idx < 0:
            print(f"未找到软件记录：{args.name}")
            sys.exit(1)
        removed = records.pop(idx)
        print(f"准备移除应用：{removed['name']}（apply={args.apply}）")
        if args.apply:
            b_dir = create_backup(excel_path, f"移除应用 {args.name}")
            if b_dir:
                print(f"已生成安全备份：{b_dir}")
            save_records(excel_path, args.device, records)
            print(f"移除成功，台账已重排保存至：{excel_path}")

    elif args.command == "sync":
        dashboard_script = os.path.join(get_documents_dir(), 'my-pc-ledger', 'sync_dashboard.py')
        if os.path.exists(dashboard_script):
            os.system(f'python "{dashboard_script}"')
            print("已刷新 my-pc-ledger/README.md 在线看板。")
        else:
            print(f"未找到看板同步脚本：{dashboard_script}")

if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    main()
